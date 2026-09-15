#!/usr/bin/env python3
"""Measure dent volumes from false-color images embedded in an Excel workbook.

The workbook is treated as read-only. Images are associated with rows by their
Excel drawing anchor, physical scale is recovered from the companion plot with
Tesseract OCR, and color is converted to depth with either an embedded color
legend or an explicit depth range.
"""

from __future__ import annotations

import argparse
import csv
import io
import math
import re
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps


@dataclass(frozen=True)
class EmbeddedImage:
    data: bytes
    image_format: str

    def open(self) -> Image.Image:
        image = Image.open(io.BytesIO(self.data))
        image.load()
        return image.convert("RGB")


@dataclass
class DentRegion:
    label: int
    mask: np.ndarray
    area_mm2: float
    volume_mm3: float
    min_depth_mm: float
    mean_depth_mm: float
    minimum_color_fraction: float
    bbox: tuple[int, int, int, int]


@dataclass(frozen=True)
class ScaleCalibration:
    mm_per_pixel_x: float
    mm_per_pixel_y: float
    x_extent_mm: float | None
    y_extent_mm: float | None
    source: str


@dataclass(frozen=True)
class LegendCalibration:
    palette_rgb: np.ndarray
    palette_depth_mm: np.ndarray
    depth_min_mm: float
    depth_max_mm: float
    source: str


def normalize_header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def resolve_column(worksheet, selector: str, search_rows: int) -> tuple[int, int]:
    """Resolve an Excel column letter or a header found near the top of a sheet."""
    if re.fullmatch(r"[A-Za-z]{1,3}", selector):
        index = column_index_from_string(selector.upper())
        return index, 0

    wanted = normalize_header(selector)
    matches: list[tuple[int, int]] = []
    for row in range(1, min(search_rows, worksheet.max_row) + 1):
        for column in range(1, worksheet.max_column + 1):
            if normalize_header(worksheet.cell(row, column).value) == wanted:
                matches.append((column, row))
    if not matches:
        raise ValueError(
            f"Column header {selector!r} was not found in the first "
            f"{search_rows} rows of sheet {worksheet.title!r}."
        )
    columns = {column for column, _ in matches}
    if len(columns) > 1:
        locations = ", ".join(
            f"row {row}, column {column}" for column, row in matches
        )
        raise ValueError(f"Column header {selector!r} is ambiguous: {locations}.")
    return matches[0]


def index_embedded_images(worksheet) -> dict[tuple[int, int], EmbeddedImage]:
    indexed: dict[tuple[int, int], EmbeddedImage] = {}
    for excel_image in worksheet._images:  # openpyxl exposes anchors only here.
        anchor = getattr(excel_image.anchor, "_from", None)
        if anchor is None:
            continue
        key = (anchor.row + 1, anchor.col + 1)
        if key in indexed:
            raise ValueError(f"More than one image is anchored to cell {key}.")
        data = excel_image._data()
        image_format = (getattr(excel_image, "format", None) or "png").lower()
        indexed[key] = EmbeddedImage(data=data, image_format=image_format)
    return indexed


def is_plain_black(image: Image.Image, nonblack_fraction: float = 0.01) -> bool:
    rgb = np.asarray(image.convert("RGB"), dtype=np.uint8)
    bright = rgb.max(axis=2) > 12
    return float(bright.mean()) < nonblack_fraction


def _largest_contiguous_run(values: np.ndarray) -> np.ndarray:
    if values.size == 0:
        return values
    cuts = np.where(np.diff(values) > 1)[0] + 1
    groups = np.split(values, cuts)
    return max(groups, key=len)


def detect_plot_bbox(image: Image.Image) -> tuple[int, int, int, int]:
    """Locate the dense colored plot inside a black-background scale image."""
    rgb = np.asarray(image.convert("RGB"), dtype=np.int16)
    chroma = (rgb.max(axis=2) - rgb.min(axis=2) > 22) & (rgb.max(axis=2) > 32)
    row_counts = chroma.sum(axis=1)
    if int(row_counts.max(initial=0)) == 0:
        raise ValueError("No colored plot was found in the scale image.")
    rows = np.where(row_counts > row_counts.max() * 0.52)[0]
    rows = _largest_contiguous_run(rows)
    if rows.size < image.height * 0.2:
        raise ValueError("The colored plot is too small to calibrate reliably.")
    y0, y1 = int(rows[0]), int(rows[-1])
    col_counts = chroma[y0 : y1 + 1].sum(axis=0)
    columns = np.where(col_counts > (y1 - y0 + 1) * 0.52)[0]
    columns = _largest_contiguous_run(columns)
    if columns.size < image.width * 0.2:
        raise ValueError("The colored plot width is too small to calibrate reliably.")
    return int(columns[0]), y0, int(columns[-1]), y1


def _image_png_bytes(image: Image.Image) -> bytes:
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def ocr_tokens(
    image: Image.Image, tesseract_command: str, page_segmentation_mode: int
) -> list[dict[str, object]]:
    command = [
        tesseract_command,
        "stdin",
        "stdout",
        "--psm",
        str(page_segmentation_mode),
        "tsv",
    ]
    try:
        completed = subprocess.run(
            command,
            input=_image_png_bytes(image),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=True,
        )
    except FileNotFoundError as exc:
        raise RuntimeError(
            f"Tesseract was not found at {tesseract_command!r}. Install it or "
            "supply explicit scale/depth calibration options."
        ) from exc
    except subprocess.CalledProcessError as exc:
        message = exc.stderr.decode("utf-8", errors="replace").strip()
        raise RuntimeError(f"Tesseract failed: {message}") from exc

    lines = completed.stdout.decode("utf-8", errors="replace").splitlines()
    if not lines:
        return []
    reader = csv.DictReader(lines, delimiter="\t")
    tokens: list[dict[str, object]] = []
    for item in reader:
        text = (item.get("text") or "").strip()
        if not text:
            continue
        try:
            tokens.append(
                {
                    "text": text,
                    "left": int(item["left"]),
                    "top": int(item["top"]),
                    "width": int(item["width"]),
                    "height": int(item["height"]),
                    "confidence": float(item["conf"]),
                }
            )
        except (KeyError, TypeError, ValueError):
            continue
    return tokens


def numeric_value(text: str) -> float | None:
    cleaned = text.strip().replace(",", ".")
    cleaned = cleaned.replace("−", "-").replace("–", "-").replace("—", "-")
    # Tesseract often interprets a short leading minus sign as a quote.
    if cleaned[:1] in {"'", '"', "`", "‘", "’", "“", "”"}:
        cleaned = "-" + cleaned[1:]
    match = re.search(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", cleaned)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def ocr_xy_extents(
    scale_image: Image.Image, tesseract_command: str
) -> tuple[float | None, float | None, tuple[int, int, int, int]]:
    bbox = detect_plot_bbox(scale_image)
    x0, y0, x1, y1 = bbox
    tokens = ocr_tokens(scale_image, tesseract_command, 11)
    numeric: list[tuple[dict[str, object], float]] = []
    for token in tokens:
        value = numeric_value(str(token["text"]))
        if value is not None and value > 0:
            numeric.append((token, value))

    plot_width = x1 - x0 + 1
    x_candidates = [
        (token, value)
        for token, value in numeric
        if int(token["top"]) > y1 + 8
        and int(token["left"]) > x0 + plot_width * 0.45
    ]
    y_candidates = [
        (token, value)
        for token, value in numeric
        if int(token["left"]) + int(token["width"]) < x0 + 5
        and abs(
            int(token["top"]) + int(token["height"]) / 2 - y0
        ) < max(35, scale_image.height * 0.13)
    ]

    x_extent = None
    if x_candidates:
        token, x_extent = max(
            x_candidates,
            key=lambda pair: (int(pair[0]["top"]), float(pair[0]["confidence"])),
        )
    y_extent = None
    if y_candidates:
        token, y_extent = min(
            y_candidates,
            key=lambda pair: (
                abs(int(pair[0]["top"]) - y0),
                -float(pair[0]["confidence"]),
            ),
        )
    return x_extent, y_extent, bbox


def calibrate_scale(
    color_image: Image.Image,
    scale_image: Image.Image,
    tesseract_command: str,
) -> ScaleCalibration:
    x_extent, y_extent, _ = ocr_xy_extents(scale_image, tesseract_command)
    if x_extent is None and y_extent is None:
        raise ValueError("OCR could not read either physical image extent.")
    mm_x = x_extent / color_image.width if x_extent is not None else None
    mm_y = y_extent / color_image.height if y_extent is not None else None
    if mm_x is None:
        mm_x = mm_y
    if mm_y is None:
        mm_y = mm_x
    assert mm_x is not None and mm_y is not None
    if not (0.001 < mm_x < 100 and 0.001 < mm_y < 100):
        raise ValueError(
            f"Implausible OCR scale: {mm_x:.6g} x {mm_y:.6g} mm/pixel."
        )
    return ScaleCalibration(mm_x, mm_y, x_extent, y_extent, "ocr_row")


def find_colorbar(
    legend_image: Image.Image,
) -> tuple[np.ndarray, tuple[int, int, int, int]]:
    rgb = np.asarray(legend_image.convert("RGB"), dtype=np.int16)
    chroma = (rgb.max(axis=2) - rgb.min(axis=2) > 35) & (rgb.max(axis=2) > 35)
    column_counts = chroma.sum(axis=0)
    center_x = int(column_counts.argmax())
    if int(column_counts[center_x]) < legend_image.height * 0.3:
        raise ValueError("No vertical color bar was found in the legend image.")
    active_rows = _largest_contiguous_run(np.where(chroma[:, center_x])[0])
    if active_rows.size < legend_image.height * 0.3:
        raise ValueError("The legend color bar is too short.")
    y0, y1 = int(active_rows[0]), int(active_rows[-1])

    active_columns = np.where(chroma[y0 : y1 + 1].sum(axis=0) > (y1 - y0) * 0.7)[0]
    local_columns = active_columns[np.abs(active_columns - center_x) < 20]
    local_columns = _largest_contiguous_run(local_columns)
    x0, x1 = int(local_columns[0]), int(local_columns[-1])
    colors = np.median(
        rgb[y0 : y1 + 1, x0 : x1 + 1], axis=1
    ).astype(np.uint8)
    return colors, (x0, y0, x1, y1)


def ocr_legend_range(
    legend_image: Image.Image,
    colorbar_bbox: tuple[int, int, int, int],
    tesseract_command: str,
    max_abs_depth_mm: float,
) -> tuple[float, float]:
    _, y0, x1, y1 = colorbar_bbox
    tokens = ocr_tokens(legend_image, tesseract_command, 6)
    values: list[tuple[dict[str, object], float]] = []
    for token in tokens:
        value = numeric_value(str(token["text"]))
        if value is not None:
            values.append((token, value))
    def vertical_center(item: tuple[dict[str, object], float]) -> float:
        token = item[0]
        return int(token["top"]) + int(token["height"]) / 2

    top_full = min(values, key=lambda item: abs(vertical_center(item) - y0)) if values else None

    def read_endpoint(crop: Image.Image) -> float | None:
        prepared = ImageOps.autocontrast(ImageOps.grayscale(crop))
        prepared = prepared.resize(
            (prepared.width * 5, prepared.height * 5), Image.Resampling.LANCZOS
        )
        candidates = []
        for token in ocr_tokens(prepared, tesseract_command, 7):
            value = numeric_value(str(token["text"]))
            if value is not None:
                candidates.append((float(token["confidence"]), value))
        if not candidates:
            return None
        return max(candidates, key=lambda item: item[0])[1]

    top_crop = legend_image.crop(
        (x1 + 3, max(0, y0 - 22), legend_image.width, min(legend_image.height, y0 + 25))
    )
    bottom_crop = legend_image.crop(
        (x1 + 3, max(0, y1 - 18), legend_image.width, legend_image.height)
    )
    top_cropped = read_endpoint(top_crop)
    bottom_value = read_endpoint(bottom_crop)

    top_value = None
    if top_full is not None and abs(top_full[1]) <= max_abs_depth_mm:
        top_value = top_full[1]
    elif top_cropped is not None and abs(top_cropped) <= max_abs_depth_mm:
        top_value = top_cropped
    if top_value is None or bottom_value is None:
        raise ValueError("OCR could not read the color legend endpoints reliably.")
    if abs(bottom_value) > max_abs_depth_mm:
        raise ValueError(
            f"Legend minimum {bottom_value:g} mm exceeds the OCR safety limit "
            f"of {max_abs_depth_mm:g} mm."
        )
    if top_value <= bottom_value:
        raise ValueError(
            f"Invalid legend range from OCR: top={top_value}, bottom={bottom_value}."
        )
    return bottom_value, top_value


def calibrate_legend(
    legend_image: Image.Image,
    tesseract_command: str,
    max_abs_depth_mm: float,
    explicit_depth_range: tuple[float, float] | None = None,
) -> LegendCalibration:
    legend_colors, legend_bbox = find_colorbar(legend_image)
    palette_rgb = resample_palette(legend_colors)
    if explicit_depth_range is None:
        depth_min_mm, depth_max_mm = ocr_legend_range(
            legend_image,
            legend_bbox,
            tesseract_command,
            max_abs_depth_mm,
        )
        source = "row_legend_ocr"
    else:
        depth_min_mm, depth_max_mm = explicit_depth_range
        source = "explicit_range_with_row_legend"
    palette_depth_mm = np.linspace(depth_max_mm, depth_min_mm, len(palette_rgb))
    return LegendCalibration(
        palette_rgb=palette_rgb,
        palette_depth_mm=palette_depth_mm,
        depth_min_mm=depth_min_mm,
        depth_max_mm=depth_max_mm,
        source=source,
    )


def generated_jet_palette(size: int = 256) -> np.ndarray:
    value = np.linspace(0.0, 1.0, size, dtype=np.float32)
    red = np.clip(1.5 - np.abs(4.0 * value - 3.0), 0.0, 1.0)
    green = np.clip(1.5 - np.abs(4.0 * value - 2.0), 0.0, 1.0)
    blue = np.clip(1.5 - np.abs(4.0 * value - 1.0), 0.0, 1.0)
    return np.rint(np.column_stack((red, green, blue)) * 255).astype(np.uint8)


def resample_palette(colors: np.ndarray, size: int = 256) -> np.ndarray:
    source = np.linspace(0.0, 1.0, len(colors))
    target = np.linspace(0.0, 1.0, size)
    channels = [np.interp(target, source, colors[:, channel]) for channel in range(3)]
    return np.rint(np.column_stack(channels)).astype(np.uint8)


def map_colors_to_depth(
    image: Image.Image,
    palette_rgb: np.ndarray,
    palette_depth_mm: np.ndarray,
    max_color_distance: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    pixels = np.asarray(image.convert("RGB"), dtype=np.uint8)
    flat = pixels.reshape(-1, 3)
    indices = np.empty(len(flat), dtype=np.int32)
    distances = np.empty(len(flat), dtype=np.float32)
    palette = palette_rgb.astype(np.int16)
    chunk_size = 8192
    for start in range(0, len(flat), chunk_size):
        stop = min(start + chunk_size, len(flat))
        difference = flat[start:stop, None, :].astype(np.int16) - palette[None, :, :]
        squared = np.sum(difference.astype(np.int32) ** 2, axis=2)
        nearest = np.argmin(squared, axis=1)
        indices[start:stop] = nearest
        distances[start:stop] = np.sqrt(squared[np.arange(stop - start), nearest])
    shape = pixels.shape[:2]
    depth = palette_depth_mm[indices].reshape(shape)
    valid = (distances <= max_color_distance).reshape(shape)
    return depth, valid, indices.reshape(shape)


def connected_components(mask: np.ndarray) -> list[np.ndarray]:
    height, width = mask.shape
    visited = np.zeros_like(mask, dtype=bool)
    components: list[np.ndarray] = []
    for seed in np.argwhere(mask):
        row, column = int(seed[0]), int(seed[1])
        if visited[row, column]:
            continue
        visited[row, column] = True
        stack = [(row, column)]
        points: list[tuple[int, int]] = []
        while stack:
            current_row, current_column = stack.pop()
            points.append((current_row, current_column))
            for row_delta in (-1, 0, 1):
                for column_delta in (-1, 0, 1):
                    if row_delta == 0 and column_delta == 0:
                        continue
                    neighbor_row = current_row + row_delta
                    neighbor_column = current_column + column_delta
                    if (
                        0 <= neighbor_row < height
                        and 0 <= neighbor_column < width
                        and mask[neighbor_row, neighbor_column]
                        and not visited[neighbor_row, neighbor_column]
                    ):
                        visited[neighbor_row, neighbor_column] = True
                        stack.append((neighbor_row, neighbor_column))
        component = np.zeros_like(mask, dtype=bool)
        rows, columns = zip(*points)
        component[np.asarray(rows), np.asarray(columns)] = True
        components.append(component)
    return components


def analyze_dents(
    depth_mm: np.ndarray,
    valid_mask: np.ndarray,
    palette_indices: np.ndarray,
    pixel_area_mm2: float,
    depth_threshold_mm: float,
    baseline_depth_mm: float,
    min_region_area_mm2: float,
    min_region_pixels: int,
    exclude_edge_regions: bool,
    edge_margin_pixels: int,
    palette_depth_mm: np.ndarray,
) -> list[DentRegion]:
    footprint = valid_mask & (depth_mm <= depth_threshold_mm)
    minimum_core_pixels = max(
        min_region_pixels,
        int(math.ceil(min_region_area_mm2 / max(pixel_area_mm2, 1e-12))),
    )
    regions: list[DentRegion] = []
    depth_floor_index = int(np.argmin(palette_depth_mm))
    for component in connected_components(footprint):
        component_pixels = int(np.count_nonzero(component))
        if component_pixels < minimum_core_pixels:
            continue
        rows, columns = np.where(component)
        touches_edge = (
            rows.min() <= edge_margin_pixels
            or columns.min() <= edge_margin_pixels
            or rows.max() >= depth_mm.shape[0] - 1 - edge_margin_pixels
            or columns.max() >= depth_mm.shape[1] - 1 - edge_margin_pixels
        )
        if exclude_edge_regions and touches_edge:
            continue
        component_depths = depth_mm[component]
        depression = np.maximum(baseline_depth_mm - component_depths, 0.0)
        minimum_color = np.abs(palette_indices[component] - depth_floor_index) <= 1
        regions.append(
            DentRegion(
                label=0,
                mask=component,
                area_mm2=float(component.sum() * pixel_area_mm2),
                volume_mm3=float(depression.sum() * pixel_area_mm2),
                min_depth_mm=float(component_depths.min()),
                mean_depth_mm=float(component_depths.mean()),
                minimum_color_fraction=float(minimum_color.mean()),
                bbox=(
                    int(columns.min()),
                    int(rows.min()),
                    int(columns.max()),
                    int(rows.max()),
                ),
            )
        )
    regions.sort(key=lambda region: region.volume_mm3, reverse=True)
    for label, region in enumerate(regions, start=1):
        region.label = label
    return regions


def _font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    names = [
        "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        if bold
        else "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    for name in names:
        try:
            return ImageFont.truetype(name, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _draw_text_lines(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    lines: Iterable[tuple[str, ImageFont.ImageFont, str]],
    spacing: int = 7,
) -> int:
    x, y = xy
    for text, font, color in lines:
        draw.text((x, y), text, font=font, fill=color)
        box = draw.textbbox((x, y), text, font=font)
        y = box[3] + spacing
    return y


def annotate_image(
    original: Image.Image,
    regions: Sequence[DentRegion],
    measurement_name: str,
    scale: ScaleCalibration,
    depth_range: tuple[float, float],
    threshold_mm: float,
    min_area_mm2: float,
    legend_is_clipped: bool,
) -> Image.Image:
    base = original.convert("RGBA")
    overlay = Image.new("RGBA", base.size, (0, 0, 0, 0))
    overlay_array = np.asarray(overlay).copy()
    colors = [(255, 181, 71), (41, 211, 255), (255, 92, 143), (154, 230, 95)]
    for index, region in enumerate(regions):
        color = colors[index % len(colors)]
        overlay_array[region.mask, :3] = color
        overlay_array[region.mask, 3] = 72
    overlay = Image.fromarray(overlay_array, mode="RGBA")
    composed = Image.alpha_composite(base, overlay)
    draw = ImageDraw.Draw(composed)
    label_font = _font(max(13, original.width // 27), bold=True)
    for index, region in enumerate(regions):
        color = colors[index % len(colors)]
        eroded = np.asarray(
            Image.fromarray((region.mask * 255).astype(np.uint8)).filter(
                ImageFilter.MinFilter(3)
            )
        ) > 0
        boundary = region.mask & ~eroded
        outline = Image.fromarray((boundary * 255).astype(np.uint8)).filter(
            ImageFilter.MaxFilter(5)
        )
        stroke = Image.new("RGBA", base.size, color + (255,))
        composed.paste(stroke, (0, 0), outline)
        draw = ImageDraw.Draw(composed)
        x0, y0, x1, y1 = region.bbox
        draw.rectangle((x0, y0, x1, y1), outline=color + (255,), width=2)
        label = f"Dent {region.label}: {region.volume_mm3:.3f} mm³"
        text_box = draw.textbbox((0, 0), label, font=label_font)
        label_width = text_box[2] - text_box[0] + 12
        label_height = text_box[3] - text_box[1] + 9
        label_x = min(max(2, x0), max(2, original.width - label_width - 2))
        label_y = max(2, y0 - label_height - 3)
        draw.rounded_rectangle(
            (label_x, label_y, label_x + label_width, label_y + label_height),
            radius=4,
            fill=(11, 18, 32, 225),
            outline=color + (255,),
        )
        draw.text((label_x + 6, label_y + 3), label, font=label_font, fill="white")

    panel_width = max(310, int(original.width * 0.80))
    canvas = Image.new("RGB", (original.width + panel_width, original.height), "#0B1220")
    canvas.paste(composed.convert("RGB"), (0, 0))
    panel = ImageDraw.Draw(canvas)
    pad = 24
    x = original.width + pad
    title_font = _font(24, bold=True)
    body_font = _font(13)
    body_bold = _font(13, bold=True)
    small_font = _font(11)
    total_volume = sum(region.volume_mm3 for region in regions)
    total_area = sum(region.area_mm2 for region in regions)
    minimum_color_pixels = sum(
        region.minimum_color_fraction * region.mask.sum() for region in regions
    )
    region_pixels = sum(region.mask.sum() for region in regions)
    minimum_color_fraction = (
        float(minimum_color_pixels / region_pixels) if region_pixels else 0.0
    )

    y = _draw_text_lines(
        panel,
        (x, 22),
        [
            ("Dent volume analysis", title_font, "#F8FAFC"),
            (measurement_name, body_font, "#A9B8CD"),
        ],
        spacing=9,
    )
    y += 8
    panel.rounded_rectangle(
        (x, y, canvas.width - pad, y + 68),
        radius=10,
        fill="#162237",
        outline="#30415C",
    )
    panel.text((x + 16, y + 10), "TOTAL VOLUME", font=small_font, fill="#8FA4BF")
    panel.text(
        (x + 16, y + 29),
        f"{total_volume:.3f} mm³",
        font=title_font,
        fill="#FFFFFF",
    )
    y += 82
    summary_lines = [
        (f"Dents retained: {len(regions)}", body_bold, "#E6EDF7"),
        (f"Total area: {total_area:.2f} mm²", body_font, "#CFD8E6"),
        (
            f"Scale: {scale.mm_per_pixel_x:.5f} × "
            f"{scale.mm_per_pixel_y:.5f} mm/px",
            body_font,
            "#CFD8E6",
        ),
        (
            f"Color depth: {depth_range[0]:.4g} to {depth_range[1]:.4g} mm",
            body_font,
            "#CFD8E6",
        ),
        (f"Dent threshold: ≤ {threshold_mm:.4g} mm", body_font, "#CFD8E6"),
        (f"Minimum region area: {min_area_mm2:.3g} mm²", body_font, "#CFD8E6"),
    ]
    y = _draw_text_lines(panel, (x, y), summary_lines, spacing=4)
    if legend_is_clipped and minimum_color_fraction > 0:
        y += 8
        warning = (
            "LOWER-BOUND VOLUME\n"
            f"{minimum_color_fraction:.1%} of dent pixels hit the legend minimum."
        )
        warning_box = (x, y, canvas.width - pad, min(canvas.height - 8, y + 54))
        panel.rounded_rectangle(warning_box, radius=8, fill="#4A2B12", outline="#F59E0B")
        panel.multiline_text(
            (x + 12, y + 9),
            warning,
            font=small_font,
            fill="#FFD58A",
            spacing=3,
        )
    elif not regions:
        y += 8
        panel.text((x, y), "No qualifying dent region found.", font=body_bold, fill="#9EE6B0")
    return canvas


def safe_filename(value: object) -> str:
    text = str(value or "unnamed").strip()
    text = re.sub(r"[^A-Za-z0-9._-]+", "_", text)
    return text.strip("._") or "unnamed"


def median_scale(calibrations: Iterable[ScaleCalibration]) -> ScaleCalibration:
    items = list(calibrations)
    return ScaleCalibration(
        mm_per_pixel_x=float(np.median([item.mm_per_pixel_x for item in items])),
        mm_per_pixel_y=float(np.median([item.mm_per_pixel_y for item in items])),
        x_extent_mm=None,
        y_extent_mm=None,
        source="median_ocr_fallback",
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Calculate dent volumes from images embedded in an Excel workbook."
    )
    parser.add_argument("input_xlsx", type=Path)
    parser.add_argument("--sheet", help="Worksheet name; defaults to the active sheet.")
    parser.add_argument("--name-column", required=True, help="Measurement-name header or letter.")
    parser.add_argument("--image-column", required=True, help="False-color image header or letter.")
    parser.add_argument("--scale-column", required=True, help="Scaled companion image header or letter.")
    parser.add_argument(
        "--legend-column",
        help="Optional color-legend header or letter. Recommended when present.",
    )
    parser.add_argument("--header-search-rows", type=int, default=20)
    parser.add_argument("--output-dir", type=Path, default=Path("analysis_output"))
    parser.add_argument("--depth-min-mm", type=float)
    parser.add_argument("--depth-max-mm", type=float)
    parser.add_argument(
        "--dent-depth-threshold-mm",
        type=float,
        default=-0.04,
        help="Pixels at or below this depth form the dent footprint (default: -0.04).",
    )
    parser.add_argument("--baseline-depth-mm", type=float, default=0.0)
    parser.add_argument("--min-region-area-mm2", type=float, default=20.0)
    parser.add_argument("--min-region-pixels", type=int, default=25)
    parser.add_argument("--mm-per-pixel", type=float, help="Override OCR with square pixels.")
    parser.add_argument("--mm-per-pixel-x", type=float)
    parser.add_argument("--mm-per-pixel-y", type=float)
    parser.add_argument("--max-color-distance", type=float, default=80.0)
    parser.add_argument(
        "--ocr-max-abs-depth-mm",
        type=float,
        default=20.0,
        help="Reject implausible OCR legend endpoints beyond this magnitude.",
    )
    parser.add_argument(
        "--legend-is-clipped",
        action="store_true",
        help=(
            "Mark volumes as lower bounds when dent pixels reach the legend minimum. "
            "Do not use this when each legend contains the real data minimum."
        ),
    )
    parser.add_argument("--include-edge-regions", action="store_true")
    parser.add_argument(
        "--edge-margin-pixels",
        type=int,
        default=5,
        help="Treat regions within this many pixels of an image edge as edge regions.",
    )
    parser.add_argument("--tesseract-command", default="tesseract")
    return parser


def validate_args(args: argparse.Namespace) -> None:
    if (args.depth_min_mm is None) != (args.depth_max_mm is None):
        raise ValueError("Supply both --depth-min-mm and --depth-max-mm, or neither.")
    if args.depth_min_mm is not None and args.depth_min_mm >= args.depth_max_mm:
        raise ValueError("--depth-min-mm must be less than --depth-max-mm.")
    if (args.mm_per_pixel_x is None) != (args.mm_per_pixel_y is None):
        raise ValueError("Supply both --mm-per-pixel-x and --mm-per-pixel-y.")
    if args.mm_per_pixel is not None and args.mm_per_pixel_x is not None:
        raise ValueError("Use either --mm-per-pixel or the X/Y pair, not both.")
    if (
        args.min_region_area_mm2 < 0
        or args.min_region_pixels < 1
        or args.edge_margin_pixels < 0
    ):
        raise ValueError("Region thresholds must be non-negative.")
    if args.baseline_depth_mm <= args.dent_depth_threshold_mm:
        raise ValueError(
            "--baseline-depth-mm must be greater than "
            "--dent-depth-threshold-mm."
        )
    if args.max_color_distance <= 0:
        raise ValueError("--max-color-distance must be positive.")
    if args.ocr_max_abs_depth_mm <= 0:
        raise ValueError("--ocr-max-abs-depth-mm must be positive.")
    for label, value in (
        ("--mm-per-pixel", args.mm_per_pixel),
        ("--mm-per-pixel-x", args.mm_per_pixel_x),
        ("--mm-per-pixel-y", args.mm_per_pixel_y),
    ):
        if value is not None and value <= 0:
            raise ValueError(f"{label} must be positive.")


def explicit_scale(args: argparse.Namespace) -> ScaleCalibration | None:
    if args.mm_per_pixel is not None:
        return ScaleCalibration(
            args.mm_per_pixel, args.mm_per_pixel, None, None, "explicit"
        )
    if args.mm_per_pixel_x is not None:
        return ScaleCalibration(
            args.mm_per_pixel_x,
            args.mm_per_pixel_y,
            None,
            None,
            "explicit_xy",
        )
    return None


def run(args: argparse.Namespace) -> Path:
    validate_args(args)
    if not args.input_xlsx.is_file():
        raise FileNotFoundError(args.input_xlsx)
    workbook = load_workbook(args.input_xlsx, read_only=False, data_only=True)
    worksheet = workbook[args.sheet] if args.sheet else workbook.active
    name_column, _ = resolve_column(worksheet, args.name_column, args.header_search_rows)
    image_column, image_header_row = resolve_column(
        worksheet, args.image_column, args.header_search_rows
    )
    scale_column, _ = resolve_column(worksheet, args.scale_column, args.header_search_rows)
    legend_column = None
    if args.legend_column:
        legend_column, _ = resolve_column(
            worksheet, args.legend_column, args.header_search_rows
        )
    images = index_embedded_images(worksheet)
    rows = sorted(
        row for row, column in images if column == image_column and row > image_header_row
    )
    if not rows:
        raise ValueError(f"No images are anchored in column {args.image_column!r}.")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    source_stem = safe_filename(args.input_xlsx.stem)

    override_scale = explicit_scale(args)
    row_scales: dict[int, ScaleCalibration] = {}
    scale_errors: dict[int, str] = {}
    usable_rows: list[int] = []
    skipped_black = 0
    for row in rows:
        color = images[(row, image_column)].open()
        if is_plain_black(color):
            skipped_black += 1
            continue
        usable_rows.append(row)
        if override_scale is not None:
            row_scales[row] = override_scale
            continue
        scale_item = images.get((row, scale_column))
        if scale_item is None:
            scale_errors[row] = "No scale image is anchored in this row."
            continue
        scale_image = scale_item.open()
        if is_plain_black(scale_image):
            scale_errors[row] = "The scale image is black."
            continue
        try:
            row_scales[row] = calibrate_scale(
                color, scale_image, args.tesseract_command
            )
        except (RuntimeError, ValueError) as exc:
            scale_errors[row] = str(exc)

    if override_scale is None and not row_scales:
        details = "; ".join(f"row {row}: {error}" for row, error in scale_errors.items())
        raise ValueError(f"No row could be scale-calibrated. {details}")
    fallback_scale = median_scale(row_scales.values()) if row_scales else override_scale
    assert fallback_scale is not None

    explicit_depth_range = (
        (args.depth_min_mm, args.depth_max_mm)
        if args.depth_min_mm is not None
        else None
    )
    row_legends: dict[int, LegendCalibration] = {}
    if legend_column is not None:
        legend_errors: dict[int, str] = {}
        for row in usable_rows:
            legend_item = images.get((row, legend_column))
            if legend_item is None:
                legend_errors[row] = "No legend image is anchored in this row."
                continue
            legend_image = legend_item.open()
            if is_plain_black(legend_image):
                legend_errors[row] = "The legend image is black."
                continue
            try:
                row_legends[row] = calibrate_legend(
                    legend_image=legend_image,
                    tesseract_command=args.tesseract_command,
                    max_abs_depth_mm=args.ocr_max_abs_depth_mm,
                    explicit_depth_range=explicit_depth_range,
                )
            except (RuntimeError, ValueError) as exc:
                legend_errors[row] = str(exc)
        if legend_errors:
            details = "; ".join(
                f"row {row}: {error}" for row, error in legend_errors.items()
            )
            raise ValueError(f"Row-specific legend calibration failed. {details}")
    else:
        if explicit_depth_range is None:
            raise ValueError(
                "Depth calibration is missing. Supply --legend-column or both "
                "--depth-min-mm and --depth-max-mm."
            )
        depth_min_mm, depth_max_mm = explicit_depth_range
        palette_rgb = generated_jet_palette()
        shared_legend = LegendCalibration(
            palette_rgb=palette_rgb,
            palette_depth_mm=np.linspace(
                depth_min_mm, depth_max_mm, len(palette_rgb)
            ),
            depth_min_mm=depth_min_mm,
            depth_max_mm=depth_max_mm,
            source="explicit_range_generated_jet",
        )
        row_legends = {row: shared_legend for row in usable_rows}

    records: list[dict[str, object]] = []
    for row in usable_rows:
        color_item = images[(row, image_column)]
        color = color_item.open()
        scale = row_scales.get(row, fallback_scale)
        scale_note = ""
        if row not in row_scales:
            scale_note = scale_errors.get(row, "Per-row scale unavailable.")
        legend = row_legends[row]
        palette_rgb = legend.palette_rgb
        palette_depth_mm = legend.palette_depth_mm
        depth_min_mm = legend.depth_min_mm
        depth_max_mm = legend.depth_max_mm
        pixel_area_mm2 = scale.mm_per_pixel_x * scale.mm_per_pixel_y
        depth, valid, palette_indices = map_colors_to_depth(
            color, palette_rgb, palette_depth_mm, args.max_color_distance
        )
        regions = analyze_dents(
            depth_mm=depth,
            valid_mask=valid,
            palette_indices=palette_indices,
            pixel_area_mm2=pixel_area_mm2,
            depth_threshold_mm=args.dent_depth_threshold_mm,
            baseline_depth_mm=args.baseline_depth_mm,
            min_region_area_mm2=args.min_region_area_mm2,
            min_region_pixels=args.min_region_pixels,
            exclude_edge_regions=not args.include_edge_regions,
            edge_margin_pixels=args.edge_margin_pixels,
            palette_depth_mm=palette_depth_mm,
        )
        measurement_value = worksheet.cell(row, name_column).value
        measurement_name = str(measurement_value or f"row_{row}")
        basename = f"{source_stem}__{safe_filename(measurement_name)}"
        original_suffix = "." + safe_filename(color_item.image_format.lower())
        if original_suffix not in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}:
            original_suffix = ".png"
        original_path = args.output_dir / f"{basename}__original{original_suffix}"
        if original_suffix == ".png" or color_item.image_format.lower() == "png":
            original_path.write_bytes(color_item.data)
        else:
            color.save(original_path)
        annotated_path = args.output_dir / f"{basename}__annotated.png"
        annotated = annotate_image(
            color,
            regions,
            measurement_name,
            scale,
            (depth_min_mm, depth_max_mm),
            args.dent_depth_threshold_mm,
            args.min_region_area_mm2,
            args.legend_is_clipped,
        )
        annotated.save(annotated_path)

        total_pixels = sum(int(region.mask.sum()) for region in regions)
        minimum_color_pixels = sum(
            region.minimum_color_fraction * int(region.mask.sum())
            for region in regions
        )
        minimum_color_fraction = (
            minimum_color_pixels / total_pixels if total_pixels else 0.0
        )
        volume_is_lower_bound = bool(
            args.legend_is_clipped and minimum_color_fraction > 0
        )
        notes = []
        if scale_note:
            notes.append(f"Scale fallback used: {scale_note}")
        if volume_is_lower_bound:
            notes.append(
                "Reported volume is a lower bound because --legend-is-clipped was "
                "set and dent pixels reach the legend minimum."
            )
        records.append(
            {
                "workbook": args.input_xlsx.name,
                "sheet": worksheet.title,
                "excel_row": row,
                "measurement_data_name": measurement_name,
                "status": "analyzed",
                "dent_count": len(regions),
                "dent_volumes_mm3": ";".join(
                    f"{region.label}:{region.volume_mm3:.6f}" for region in regions
                ),
                "dent_areas_mm2": ";".join(
                    f"{region.label}:{region.area_mm2:.6f}" for region in regions
                ),
                "total_volume_mm3": round(
                    sum(region.volume_mm3 for region in regions), 6
                ),
                "total_area_mm2": round(sum(region.area_mm2 for region in regions), 6),
                "minimum_depth_mm": (
                    round(min(region.min_depth_mm for region in regions), 6)
                    if regions
                    else ""
                ),
                "mean_dent_depth_mm": (
                    round(
                        float(
                            np.average(
                                [region.mean_depth_mm for region in regions],
                                weights=[region.mask.sum() for region in regions],
                            )
                        ),
                        6,
                    )
                    if regions
                    else ""
                ),
                "mm_per_pixel_x": round(scale.mm_per_pixel_x, 8),
                "mm_per_pixel_y": round(scale.mm_per_pixel_y, 8),
                "scale_source": scale.source,
                "legend_source": legend.source,
                "depth_min_mm": depth_min_mm,
                "depth_max_mm": depth_max_mm,
                "dent_depth_threshold_mm": args.dent_depth_threshold_mm,
                "baseline_depth_mm": args.baseline_depth_mm,
                "minimum_region_area_mm2": args.min_region_area_mm2,
                "valid_color_fraction": round(float(valid.mean()), 6),
                "minimum_legend_color_fraction": round(
                    minimum_color_fraction, 6
                ),
                "legend_is_clipped": args.legend_is_clipped,
                "volume_is_lower_bound": volume_is_lower_bound,
                "original_image": original_path.name,
                "annotated_image": annotated_path.name,
                "notes": " ".join(notes),
            }
        )
        bound = " (lower bound)" if volume_is_lower_bound else ""
        total_volume = sum(region.volume_mm3 for region in regions)
        print(
            f"row {row:>3} {measurement_name:<24} "
            f"dents={len(regions):>2} volume={total_volume:.3f} mm^3{bound}"
        )

    csv_path = args.output_dir / f"{source_stem}__dent_volumes.csv"
    if records:
        with csv_path.open("w", newline="", encoding="utf-8") as output:
            writer = csv.DictWriter(output, fieldnames=list(records[0]))
            writer.writeheader()
            writer.writerows(records)
    else:
        raise ValueError("All source images were black; no rows were analyzed.")
    print(
        f"Wrote {len(records)} analyzed rows to {csv_path}; "
        f"ignored {skipped_black} black rows."
    )
    return csv_path


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        run(args)
    except Exception as exc:  # Give CLI users one concise, actionable error.
        print(f"error: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
